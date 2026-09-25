"""Leitura das planilhas de rateio das usinas.

Estrutura de pastas (raiz do projeto):
  rateio/documents/<GERADORA>/  documentos anexados na conclusão do rateio
  rateio/lista/<AAAA-MM>/       planilhas do mês: rateio_<GERADORA>_<AAAA>_<MM>_fator<N>.xlsx

Da planilha só a aba "Distribuição por Usina" define o rateio: um bloco por
usina, com a linha-título "ENERGIA A 1  (GDII)  · Capacidade: ...", o cabeçalho
"Nova UC | Cliente | Situação | ... | % na Usina | kWh" e a linha "Total <usina>".
A linha "(bateria)" (sobra da usina) vai para a UC do endereço-bateria listado
na aba "Configuração MCP".

A UC da usina (geradora) é procurada, nesta ordem: na linha-título ou nas linhas
entre ela e o cabeçalho; numa coluna "UC Usina"/"UC Geradora" do bloco; numa
coluna de UC do snapshot de usinas da aba "Configuração MCP".
"""
import os
import re
from dataclasses import dataclass, field

import openpyxl

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PASTA_RATEIO = os.path.join(PROJECT_DIR, "rateio")
PASTA_LISTA = os.path.join(PASTA_RATEIO, "lista")
PASTA_DOCUMENTOS = os.path.join(PASTA_RATEIO, "documents")

ABA_DISTRIBUICAO = "Distribuição por Usina"
ABA_CONFIGURACAO = "Configuração MCP"
MARCADOR_BATERIA = "(bateria)"

RE_UC = re.compile(r"\d{2}/\d{7}-\d")
RE_MES = re.compile(r"^\d{4}-\d{2}$")
RE_TITULO_USINA = re.compile(r"^(?P<nome>.+?)\s*\((?P<classe>GD\s*I{1,3})\)(?:.*?Capacidade:\s*(?P<cap>[\d.,]+))?")
RE_NOME_ARQUIVO = re.compile(r"^rateio_(?P<geradora>.+?)_(?P<ano>\d{4})_(?P<mes>\d{2})", re.IGNORECASE)


@dataclass
class Beneficiaria:
    uc: str
    cliente: str
    situacao: str
    percentual: float
    kwh: float
    bateria: bool = False


@dataclass
class Usina:
    nome: str
    classe: str
    capacidade_kwh: float
    uc_geradora: str = ""
    beneficiarias: list = field(default_factory=list)

    @property
    def soma_percentual(self):
        return round(sum(b.percentual or 0 for b in self.beneficiarias), 4)

    @property
    def total_kwh(self):
        return sum(b.kwh or 0 for b in self.beneficiarias)


@dataclass
class PlanilhaRateio:
    arquivo: str
    geradora: str
    mes_referencia: str
    usinas: list = field(default_factory=list)
    uc_bateria: str = ""
    documentos: list = field(default_factory=list)
    avisos: list = field(default_factory=list)

    @property
    def total_linhas(self):
        return sum(len(u.beneficiarias) for u in self.usinas)

    @property
    def ucs_distintas(self):
        return len({b.uc for u in self.usinas for b in u.beneficiarias})


# ==================== PASTAS ====================

def listar_meses():
    """Pastas AAAA-MM de rateio/lista, da mais recente para a mais antiga."""
    if not os.path.isdir(PASTA_LISTA):
        return []
    return sorted((d for d in os.listdir(PASTA_LISTA)
                   if RE_MES.match(d) and os.path.isdir(os.path.join(PASTA_LISTA, d))), reverse=True)


def listar_planilhas(mes):
    """Caminhos das .xlsx da pasta do mês (ignora os temporários "~$" do Excel)."""
    pasta = os.path.join(PASTA_LISTA, mes)
    if not os.path.isdir(pasta):
        return []
    return [os.path.join(pasta, a) for a in sorted(os.listdir(pasta))
            if a.lower().endswith(".xlsx") and not a.startswith("~$")]


def listar_documentos(geradora):
    """Arquivos de rateio/documents/<GERADORA>, em ordem de nome (01-, 02-, ...)."""
    pasta = os.path.join(PASTA_DOCUMENTOS, geradora)
    if not os.path.isdir(pasta):
        return []
    return [os.path.join(pasta, a) for a in sorted(os.listdir(pasta))
            if os.path.isfile(os.path.join(pasta, a))]


# ==================== LEITURA ====================

def _texto(valor):
    return str(valor).strip() if valor is not None else ""


def _numero(valor):
    if isinstance(valor, (int, float)):
        return float(valor)
    try:
        return float(_texto(valor).replace("%", "").replace(",", "."))
    except ValueError:
        return None


def _uc_em(linha):
    """Primeira UC (NN/NNNNNNN-N) encontrada nas células da linha."""
    for valor in linha:
        achou = RE_UC.search(_texto(valor))
        if achou:
            return achou.group(0)
    return ""


def _indice(cabecalho, *nomes):
    """Índice da 1ª coluna cujo título contém todos os trechos de um dos nomes."""
    titulos = [_texto(c).lower() for c in cabecalho]
    for nome in nomes:
        trechos = nome.lower().split("|")
        for i, titulo in enumerate(titulos):
            if all(t in titulo for t in trechos):
                return i
    return None


def _ler_configuracao(wb):
    """Da aba "Configuração MCP": UC do endereço-bateria e, se houver coluna de UC
    no snapshot de usinas, o mapa nome da usina → UC geradora."""
    if ABA_CONFIGURACAO not in wb.sheetnames:
        return "", {}
    uc_bateria, ucs_usinas = "", {}
    secao, cabecalho = None, None
    for linha in wb[ABA_CONFIGURACAO].iter_rows(values_only=True):
        primeira = _texto(linha[0]) if linha else ""
        if primeira.startswith("Usinas ("):
            secao, cabecalho = "usinas", None
            continue
        if primeira.startswith("Endereços-Bateria"):
            secao, cabecalho = "bateria", None
            continue
        if not any(v is not None for v in linha):
            if cabecalho is not None:
                secao = None
            continue
        if secao and cabecalho is None:
            cabecalho = linha
            continue
        if secao == "bateria" and not uc_bateria:
            uc_bateria = _uc_em(linha)
        elif secao == "usinas":
            i_nome = _indice(cabecalho, "nome")
            i_uc = _indice(cabecalho, "uc", "unidade|consumidora")
            if i_nome is not None and i_uc is not None:
                uc = _uc_em([linha[i_uc]])
                if uc:
                    ucs_usinas[_texto(linha[i_nome])] = uc
    return uc_bateria, ucs_usinas


def _ler_distribuicao(ws):
    """Blocos de usina da aba "Distribuição por Usina"."""
    usinas, atual, colunas = [], None, None
    for linha in ws.iter_rows(min_row=2, values_only=True):
        primeira = _texto(linha[0]) if linha else ""
        titulo = RE_TITULO_USINA.match(primeira)
        if titulo:
            capacidade = _numero((titulo.group("cap") or "0").replace(",", "")) or 0
            atual = Usina(nome=titulo.group("nome").strip(),
                          classe=titulo.group("classe").replace(" ", ""),
                          capacidade_kwh=capacidade, uc_geradora=_uc_em(linha))
            usinas.append(atual)
            colunas = None
            continue
        if atual is None:
            continue
        if primeira.lower() == "nova uc":
            colunas = {
                "uc": 0,
                "cliente": _indice(linha, "cliente"),
                "situacao": _indice(linha, "situação", "situacao"),
                "percentual": _indice(linha, "% na usina", "%"),
                "kwh": _indice(linha, "kwh"),
                "uc_usina": _indice(linha, "uc|usina", "uc|geradora"),
            }
            continue
        if colunas is None:
            # Linhas entre o título e o cabeçalho: podem trazer a UC da usina.
            if not atual.uc_geradora:
                atual.uc_geradora = _uc_em(linha)
            continue
        if not primeira or primeira.lower().startswith("total"):
            if primeira.lower().startswith("total"):
                atual, colunas = None, None
            continue

        def celula(chave):
            i = colunas.get(chave)
            return linha[i] if i is not None and i < len(linha) else None

        if colunas["uc_usina"] is not None and not atual.uc_geradora:
            atual.uc_geradora = _uc_em([celula("uc_usina")])
        atual.beneficiarias.append(Beneficiaria(
            uc=primeira,
            cliente=_texto(celula("cliente")),
            situacao=_texto(celula("situacao")),
            percentual=_numero(celula("percentual")),
            kwh=_numero(celula("kwh")),
            bateria=primeira.lower() == MARCADOR_BATERIA,
        ))
    return usinas


def _geradora_e_mes(wb, caminho):
    """Geradora e mês pelo título da aba ("Distribuição por Usina — ENERGIA A — 2026-09")
    ou, na falta dele, pelo nome do arquivo."""
    titulo = _texto(wb[ABA_DISTRIBUICAO]["A1"].value)
    partes = [p.strip() for p in re.split(r"\s+[—-]\s+", titulo)]
    if len(partes) >= 3 and RE_MES.match(partes[-1]):
        return partes[1], partes[-1]
    nome = RE_NOME_ARQUIVO.match(os.path.basename(caminho))
    if nome:
        return nome.group("geradora").replace("_", " "), f"{nome.group('ano')}-{nome.group('mes')}"
    return "", ""


def _validar(planilha):
    avisos = planilha.avisos
    if not planilha.usinas:
        avisos.append(f"Nenhum bloco de usina encontrado na aba '{ABA_DISTRIBUICAO}'.")
    for usina in planilha.usinas:
        if not usina.uc_geradora:
            avisos.append(f"{usina.nome}: UC da usina não informada na planilha.")
        if not usina.beneficiarias:
            avisos.append(f"{usina.nome}: nenhuma beneficiária.")
        if abs(usina.soma_percentual - 100) > 0.01:
            avisos.append(f"{usina.nome}: soma dos percentuais = {usina.soma_percentual:g}% (esperado 100%).")
        vistas = set()
        for b in usina.beneficiarias:
            if b.percentual is None or b.percentual <= 0:
                avisos.append(f"{usina.nome}: percentual inválido para {b.uc} ({b.percentual}).")
            if b.bateria and not RE_UC.fullmatch(b.uc):
                avisos.append(f"{usina.nome}: linha de bateria sem UC (bateria não encontrada na aba '{ABA_CONFIGURACAO}').")
            elif not RE_UC.fullmatch(b.uc):
                avisos.append(f"{usina.nome}: UC fora do padrão NN/NNNNNNN-N: '{b.uc}'.")
            if b.uc in vistas:
                avisos.append(f"{usina.nome}: UC {b.uc} repetida na mesma usina.")
            vistas.add(b.uc)
    if planilha.geradora and not planilha.documentos:
        avisos.append(f"Sem documentos em rateio/documents/{planilha.geradora}/.")


def ler_planilha(caminho):
    """Lê e valida uma planilha de rateio.

    Returns:
        PlanilhaRateio (problemas encontrados ficam em .avisos).

    Raises:
        ValueError: se a planilha não tem a aba "Distribuição por Usina".
    """
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=False)
    if ABA_DISTRIBUICAO not in wb.sheetnames:
        raise ValueError(f"{os.path.basename(caminho)}: aba '{ABA_DISTRIBUICAO}' não encontrada "
                         f"(abas: {', '.join(wb.sheetnames)})")

    geradora, mes = _geradora_e_mes(wb, caminho)
    uc_bateria, ucs_usinas = _ler_configuracao(wb)
    usinas = _ler_distribuicao(wb[ABA_DISTRIBUICAO])
    for usina in usinas:
        if not usina.uc_geradora:
            usina.uc_geradora = ucs_usinas.get(usina.nome, "")
        for b in usina.beneficiarias:
            if b.bateria and uc_bateria:
                b.uc = uc_bateria

    planilha = PlanilhaRateio(arquivo=caminho, geradora=geradora, mes_referencia=mes,
                              usinas=usinas, uc_bateria=uc_bateria,
                              documentos=listar_documentos(geradora))
    _validar(planilha)
    return planilha
